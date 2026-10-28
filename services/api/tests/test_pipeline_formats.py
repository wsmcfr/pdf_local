from docx import Document
from openpyxl import Workbook

from app.db import connect, init_db
from app.index.vector_store import VectorStore
from app.ingest.pipeline import IngestPipeline
from app.ingest.scanner import ScannedFile
from app.utils.hashing import sha256_file


def scanned(path):
    stat = path.stat()
    return ScannedFile(
        path=path,
        root=path.parent,
        content_hash=sha256_file(path),
        size_bytes=stat.st_size,
        mtime=stat.st_mtime,
    )


def make_pipeline(tmp_path):
    conn = connect(tmp_path / "app.db")
    init_db(conn)
    store = VectorStore(tmp_path / "vectors.faiss")
    return conn, store, IngestPipeline(conn, vector_store=store)


def test_docx_is_stored_with_pages_blocks_and_chunks(tmp_path):
    path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("第一段内容")
    document.add_paragraph("第二段内容")
    document.save(path)

    conn, _store, pipeline = make_pipeline(tmp_path)
    assert pipeline._ingest_file(scanned(path)) is True

    row = conn.execute("SELECT * FROM documents WHERE path = ?", (str(path),)).fetchone()
    assert row["status"] == "ready"
    assert row["mime"].endswith("wordprocessingml.document")
    assert row["editable"] == 1
    assert row["source_format"] == "docx"
    assert row["page_count"] == 1

    blocks = conn.execute(
        "SELECT * FROM blocks WHERE document_id = ? ORDER BY order_index",
        (row["id"],),
    ).fetchall()
    assert [block["text"] for block in blocks] == ["第一段内容", "第二段内容"]

    chunks = conn.execute(
        "SELECT COUNT(*) AS n FROM chunks WHERE document_id = ?", (row["id"],)
    ).fetchone()
    assert chunks["n"] >= 1
    conn.close()


def test_xlsx_is_read_only_and_searchable(tmp_path):
    path = tmp_path / "scores.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "张三"
    sheet["B1"] = 90
    workbook.save(path)

    conn, _store, pipeline = make_pipeline(tmp_path)
    assert pipeline._ingest_file(scanned(path)) is True

    row = conn.execute("SELECT * FROM documents WHERE path = ?", (str(path),)).fetchone()
    assert row["status"] == "ready"
    assert row["editable"] == 0
    assert row["source_format"] == "xlsx"
    assert row["mime"].endswith("spreadsheetml.sheet")

    blocks = conn.execute(
        "SELECT * FROM blocks WHERE document_id = ?", (row["id"],)
    ).fetchall()
    assert [block["text"] for block in blocks] == ["张三", "90"]
    conn.close()
