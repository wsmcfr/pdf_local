from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.ingest.formats import parse_office_file


def test_docx_paragraphs_become_editable_blocks(tmp_path):
    path = tmp_path / "a.docx"
    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    document.save(path)

    parsed = parse_office_file(path)

    assert [block.text for block in parsed.blocks] == ["第一段", "第二段"]
    assert parsed.blocks[0].locator == {"paragraph": 0}
    assert parsed.blocks[0].editable is True
    assert parsed.pages[0].text == "第一段\n\n第二段"


def test_docx_tables_are_kept_as_cell_blocks(tmp_path):
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "分数"
    document.save(path)

    parsed = parse_office_file(path)

    locators = [block.locator for block in parsed.blocks]
    assert {"table": 0, "row": 0, "col": 0} in locators
    assert {"table": 0, "row": 0, "col": 1} in locators


def test_pptx_blocks_carry_slide_shape_and_bbox(tmp_path):
    path = tmp_path / "a.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "标题"
    slide.placeholders[1].text = "正文"
    presentation.save(path)

    parsed = parse_office_file(path)

    title = next(block for block in parsed.blocks if block.text == "标题")
    assert title.page_no == 1
    assert title.locator["slide"] == 1
    assert title.locator["paragraph"] == 0
    assert title.bbox is not None
    assert 0 <= title.bbox["x"] <= 1
    assert title.editable is True


def test_xlsx_cells_are_read_only(tmp_path):
    path = tmp_path / "a.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成绩"
    sheet["A1"] = "姓名"
    sheet["B1"] = "分数"
    sheet["A2"] = "张三"
    sheet["B2"] = 90
    workbook.save(path)

    parsed = parse_office_file(path)

    assert parsed.pages[0].page_no == 1
    assert "姓名 | 分数" in parsed.pages[0].text
    first = parsed.blocks[0]
    assert first.kind == "cell"
    assert first.locator == {"sheet": "成绩", "row": 1, "col": 1}
    assert first.editable is False


def test_text_file_blocks_keep_offsets(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("第一行\n第二行\n", encoding="utf-8")

    parsed = parse_office_file(path)

    assert [block.text for block in parsed.blocks] == ["第一行", "第二行"]
    assert parsed.blocks[0].locator == {"offset": 0}
    assert parsed.blocks[1].locator == {"offset": 4}
    assert parsed.pages[0].text == "第一行\n第二行\n"
